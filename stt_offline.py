
import json
import os
import time

import sounddevice as sd
from openpyxl import Workbook, load_workbook
from vosk import Model, KaldiRecognizer


# ---------------- Configuration ----------------
MODEL_PATH = "vosk-model"
SAMPLE_RATE = 16000
RECORD_SECONDS = 5
EXCEL_FILE = "latency_report.xlsx"


def record_audio():
    """Record microphone audio as mono, 16-bit PCM."""
    print(f"Speak now for {RECORD_SECONDS} seconds...")
    audio = sd.rec(
        int(RECORD_SECONDS * SAMPLE_RATE),
        samplerate=SAMPLE_RATE,
        channels=1,
        dtype="int16",
    )
    sd.wait()
    print("Recording completed.")
    return audio


def transcribe_audio(model, audio):
    """Transcribe recorded audio and measure processing time."""
    recognizer = KaldiRecognizer(model, SAMPLE_RATE)

    start = time.perf_counter()
    recognizer.AcceptWaveform(audio.tobytes())
    result = json.loads(recognizer.FinalResult())
    elapsed = time.perf_counter() - start

    return result.get("text", ""), elapsed


def save_report(text, model_seconds, recording_seconds, transcription_seconds):
    """Save the latest result and preserve all trial results."""
    total_seconds = recording_seconds + transcription_seconds
    rtf = transcription_seconds / recording_seconds

    if os.path.exists(EXCEL_FILE):
        workbook = load_workbook(EXCEL_FILE)
    else:
        workbook = Workbook()

    # Latest Run sheet
    if "Latest Run" in workbook.sheetnames:
        summary = workbook["Latest Run"]
        summary.delete_rows(1, summary.max_row)
    else:
        summary = workbook.active
        summary.title = "Latest Run"

    summary.append(["Metric", "Value", "Meaning"])
    summary.append([
        "Recognized text",
        text or "(No speech recognized)",
        "Text returned by the Vosk model",
    ])
    summary.append([
        "Model loading time",
        f"{model_seconds:.3f} seconds",
        "Time needed to load the local model",
    ])
    summary.append([
        "Recording time",
        f"{recording_seconds:.3f} seconds",
        "Duration of microphone recording",
    ])
    summary.append([
        "Transcription time",
        f"{transcription_seconds:.3f} seconds",
        "Time spent converting recorded audio into text",
    ])
    summary.append([
        "Recording + transcription",
        f"{total_seconds:.3f} seconds",
        "Recording duration plus transcription processing time",
    ])
    summary.append([
        "Real-time factor (RTF)",
        round(rtf, 3),
        "Transcription time divided by recording duration; below 1 means faster than real time",
    ])
    summary.append([
        "Offline verification",
        "Not verified",
        "Verify by running successfully with Wi-Fi disconnected",
    ])

    # All Trials sheet
    if "All Trials" in workbook.sheetnames:
        trials = workbook["All Trials"]
    else:
        trials = workbook.create_sheet("All Trials")
        trials.append([
            "Test number",
            "Recognized text",
            "Model loading (s)",
            "Recording (s)",
            "Transcription (s)",
            "Recording + transcription (s)",
            "Real-time factor (RTF)",
        ])

    test_number = trials.max_row
    trials.append([
        test_number,
        text,
        round(model_seconds, 3),
        round(recording_seconds, 3),
        round(transcription_seconds, 3),
        round(total_seconds, 3),
        round(rtf, 3),
    ])

    times = [
        trials.cell(row=row, column=5).value
        for row in range(2, trials.max_row + 1)
        if isinstance(trials.cell(row=row, column=5).value, (int, float))
    ]
    average = sum(times) / len(times) if times else 0.0

    summary.append([
        "Average transcription time",
        f"{average:.3f} seconds",
        "Average across all saved trials",
    ])

    # Format the Excel sheets
    for sheet in (summary, trials):
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions

        for column in sheet.columns:
            letter = column[0].column_letter
            longest = max(len(str(cell.value or "")) for cell in column)
            sheet.column_dimensions[letter].width = min(
                max(longest + 2, 15), 55
            )

   
    try:
        workbook.save(EXCEL_FILE)
    except PermissionError:
        backup_file = "latency_report_backup.xlsx"
        try:
            workbook.save(backup_file)
            print(f"Main report is locked. Saved report to: {backup_file}")
        except PermissionError:
            print("Both report files are locked. Close them in Excel and retry.")
            raise SystemExit(1)


    return test_number, average


def main():
    # Check the local model folder
    if not os.path.isdir(MODEL_PATH):
        print(f"ERROR: Model folder not found: {MODEL_PATH}")
        print("Place the extracted Vosk model folder beside this script.")
        raise SystemExit(1)

    # Load model once
    print("Loading Vosk model...")
    start = time.perf_counter()

    try:
        model = Model(MODEL_PATH)
    except Exception as error:
        print(f"ERROR loading Vosk model: {error}")
        raise SystemExit(1)

    model_seconds = time.perf_counter() - start
    print(f"Model loaded in {model_seconds:.3f} seconds.")

    # Record audio
    input("\nPress Enter to start recording...")

    try:
        audio = record_audio()
    except Exception as error:
        print(f"ERROR recording microphone audio: {error}")
        raise SystemExit(1)

    # Transcribe and calculate latency
    text, transcription_seconds = transcribe_audio(model, audio)
    recording_seconds = float(RECORD_SECONDS)
    total_seconds = recording_seconds + transcription_seconds
    rtf = transcription_seconds / recording_seconds

    print("\n========== SPEECH-TO-TEXT RESULT ==========")
    print("Recognized text:", text or "(No speech recognized)")
    print(f"Model loading time: {model_seconds:.3f} seconds")
    print(f"Recording time: {recording_seconds:.3f} seconds")
    print(f"Transcription time: {transcription_seconds:.3f} seconds")
    print(f"Recording + transcription: {total_seconds:.3f} seconds")
    print(f"Real-time factor (RTF): {rtf:.3f}")

    # Save Excel report
    test_number, average = save_report(
        text,
        model_seconds,
        recording_seconds,
        transcription_seconds,
    )

    print(f"\nExcel report saved to: {EXCEL_FILE}")
    print(f"Test number: {test_number}")
    print(f"Average transcription time: {average:.3f} seconds")


if __name__ == "__main__":
    main()
